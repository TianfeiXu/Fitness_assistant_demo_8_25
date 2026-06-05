from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

# Create a presentation
prs = Presentation()

# Define slide layout (title + content)
title_slide_layout = prs.slide_layouts[0]
content_slide_layout = prs.slide_layouts[1]

# Scene 1 – Opening
slide1 = prs.slides.add_slide(title_slide_layout)
title1 = slide1.shapes.title
subtitle1 = slide1.placeholders[1]
title1.text = "AI Fitness Log Assistant"
subtitle1.text = "By Tianfei Xu\nMaster of Computer Science, Colorado School of Mines"

# Scene 2 – Purpose & Tools
slide2 = prs.slides.add_slide(content_slide_layout)
title2 = slide2.shapes.title
content2 = slide2.placeholders[1]
title2.text = "Purpose & Tools"
content2.text = (
    "• Purpose: build an AI-powered fitness assistant\n"
    "• Python + Streamlit for the interface\n"
    "• Pandas for data processing\n"
    "• Altair for visualization\n"
    "• OpenAI API for parsing logs and generating suggestions\n"
    "• Demonstrates ability to learn and integrate new tools"
)

# Scene 3 – High-level Method
slide3 = prs.slides.add_slide(content_slide_layout)
title3 = slide3.shapes.title
content3 = slide3.placeholders[1]
title3.text = "High-level Method"
content3.text = (
    "1. User enters workout text\n"
    "2. System sends text to OpenAI API\n"
    "3. Parse into structured data (exercise, body part, sets, reps, minutes)\n"
    "4. Calculate calories, 7-day trends, and body-part heat map\n"
    "5. Generate personalized training suggestions"
)

# Save the presentation
file_path = "AI_Fitness_Demo_Slides_Scene1-3.pptx"
prs.save(file_path)

file_path
